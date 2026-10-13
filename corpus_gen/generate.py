"""
corpus_gen/generate.py
Generator for the 20 Novacure Labs documents in PLAN.md section 6.2.
Produces realistic content, sidecar JSON metadata, and renders files to:
  - data/raw/sop/*.pdf
  - data/raw/pdf/*.pdf
  - data/raw/sharepoint/*.docx
  - data/raw/email/*.eml
  - data/raw/teams/*.json
Supports --offline, --force, and --only DOC_ID flags.
"""

import argparse
import email.policy
from email.message import EmailMessage
import json
import os
import sys
import time
import xml.sax.saxutils
from pathlib import Path
from typing import Any, Dict, List, Optional

# Do NOT import from app/ or modules/
from corpus_gen.manifest import MANIFEST
from corpus_gen.templates import TEMPLATES
from corpus_gen.client import generate as llm_generate

# Base directory for raw files
DATA_RAW_DIR = Path("data/raw")


def sanitize_text(text: str) -> str:
    """Remove forbidden strings like 'HPLC' from any text."""
    if not text:
        return ""
    return text.replace("HPLC", "")


def ensure_fixed_sentences(text: str, fixed_sentences: List[str]) -> str:
    """Ensure every fixed sentence appears verbatim in the text.
    If missing, append it as its own distinct paragraph.
    """
    clean_text = sanitize_text(text)
    for sentence in fixed_sentences:
        clean_sentence = sanitize_text(sentence)
        if clean_sentence not in clean_text:
            clean_text = f"{clean_text}\n\n{clean_sentence}"
    return clean_text.strip()


def render_pdf(filepath: Path, spec: Dict[str, Any], content: Any) -> None:
    """Render PDF document using reportlab with clean extractable text."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    filepath.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(filepath),
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontSize=16,
        leading=20,
        spaceAfter=6,
        textColor=colors.HexColor("#1A202C"),
    )
    meta_style = ParagraphStyle(
        "DocMeta",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#4A5568"),
        spaceAfter=14,
    )
    section_title_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontSize=12,
        leading=16,
        spaceBefore=10,
        spaceAfter=4,
        textColor=colors.HexColor("#2D3748"),
    )
    body_style = ParagraphStyle(
        "DocBody",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        spaceAfter=8,
        textColor=colors.HexColor("#2D3748"),
    )

    story = []

    # Title & Metadata Header
    safe_title = xml.sax.saxutils.escape(sanitize_text(spec["title"]))
    story.append(Paragraph(f"<b>{safe_title}</b>", title_style))

    version_str = f"v{spec['version']}" if spec.get("version") else "N/A"
    meta_text = (
        f"<b>Document ID:</b> {spec['doc_id']} | <b>Version:</b> {version_str} | "
        f"<b>Effective Date:</b> {spec['date']} | <b>Author:</b> {spec['author']}"
    )
    story.append(Paragraph(meta_text, meta_style))
    story.append(Spacer(1, 10))

    # Parse sections from template or generated text
    if isinstance(content, dict) and "sections" in content:
        for heading, body in content["sections"]:
            clean_h = xml.sax.saxutils.escape(sanitize_text(heading))
            clean_b = xml.sax.saxutils.escape(sanitize_text(body))
            story.append(Paragraph(f"<b>{clean_h}</b>", section_title_style))
            story.append(Paragraph(clean_b, body_style))
    elif isinstance(content, str):
        # Paragraph-based text
        paragraphs = [p.strip() for p in content.split("\n\n") if p.strip()]
        for p in paragraphs:
            clean_p = xml.sax.saxutils.escape(sanitize_text(p))
            story.append(Paragraph(clean_p, body_style))

    doc.build(story)


def render_docx(filepath: Path, spec: Dict[str, Any], content: Any) -> None:
    """Render SharePoint document using python-docx."""
    import docx

    filepath.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()

    # Title & Subtitle
    safe_title = sanitize_text(spec["title"])
    doc.add_heading(safe_title, level=1)

    version_str = f"v{spec['version']}" if spec.get("version") else "N/A"
    doc.add_paragraph(
        f"Document ID: {spec['doc_id']} | Version: {version_str} | "
        f"Date: {spec['date']} | Author: {spec['author']} | Owner: {spec['owner']}"
    )

    if isinstance(content, dict) and "sections" in content:
        for heading, body in content["sections"]:
            doc.add_heading(sanitize_text(heading), level=2)
            doc.add_paragraph(sanitize_text(body))
    elif isinstance(content, str):
        for p in content.split("\n\n"):
            p_clean = sanitize_text(p.strip())
            if p_clean:
                if p_clean.startswith("#"):
                    doc.add_heading(p_clean.lstrip("#").strip(), level=2)
                else:
                    doc.add_paragraph(p_clean)

    doc.save(str(filepath))


def render_eml(filepath: Path, spec: Dict[str, Any], content: Any) -> None:
    """Render email thread as .eml using stdlib email package."""
    filepath.parent.mkdir(parents=True, exist_ok=True)
    msg = EmailMessage()

    from_addr = f"{spec['author'].lower().replace(' ', '.').replace('dr..', '')}@novacure.com"
    subject = spec["title"]
    date_str = f"{spec['date']} 09:00:00 +0000"

    body_text = ""
    if isinstance(content, dict):
        from_addr = content.get("from", from_addr)
        to_addr = content.get("to", "team@novacure.com")
        subject = content.get("subject", subject)
        date_str = content.get("date", date_str)
        body_text = content.get("body", "")
    elif isinstance(content, str):
        body_text = content
        to_addr = "qa.team@novacure.com"
    else:
        to_addr = "team@novacure.com"

    body_text = sanitize_text(body_text)

    msg["From"] = from_addr
    msg["To"] = to_addr
    msg["Subject"] = sanitize_text(subject)
    msg["Date"] = date_str
    msg.set_content(body_text)

    with open(filepath, "wb") as f:
        f.write(msg.as_bytes())


def render_teams(filepath: Path, spec: Dict[str, Any], content: Any) -> None:
    """Render Teams channel messages as a JSON list of message objects."""
    filepath.parent.mkdir(parents=True, exist_ok=True)
    messages: List[Dict[str, Any]] = []

    if isinstance(content, list):
        for m in content:
            messages.append({
                "msg_id": int(m["msg_id"]),
                "author": str(m["author"]),
                "timestamp": str(m["timestamp"]),
                "text": sanitize_text(str(m["text"])),
            })
    else:
        # Fallback if content was string
        lines = [line.strip() for line in str(content).split("\n") if line.strip()]
        for idx, line in enumerate(lines[:15], 1):
            messages.append({
                "msg_id": idx,
                "author": spec["author"],
                "timestamp": f"{spec['date']}T{9 + (idx % 8):02d}:00:00Z",
                "text": sanitize_text(line),
            })

    # Ensure fixed sentences appear in at least one message
    for sentence in spec.get("fixed_sentences", []):
        clean_sentence = sanitize_text(sentence)
        found = any(clean_sentence in m["text"] for m in messages)
        if not found:
            # Append fixed sentence as a message from primary author
            msg_id = len(messages) + 1
            author = spec["author"]
            if spec["doc_id"] == "TEAMS-COLDCHAIN":
                author = "Arjun Menon"
            elif spec["doc_id"] in ["TEAMS-QAHELP", "TEAMS-DEVIATIONS"]:
                author = "Priya Rao"
            messages.append({
                "msg_id": msg_id,
                "author": author,
                "timestamp": f"{spec['date']}T12:00:00Z",
                "text": clean_sentence,
            })

    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(messages, f, indent=2, ensure_ascii=False)


def write_sidecar(filepath: Path, spec: Dict[str, Any]) -> None:
    """Write <file>.meta.json sidecar file."""
    sidecar_path = Path(f"{filepath}.meta.json")
    sidecar_data = {
        "doc_id": spec["doc_id"],
        "title": spec["title"],
        "source_type": spec["source_type"],
        "version": spec["version"],
        "date": spec["date"],
        "author": spec["author"],
        "owner": spec["owner"],
        "status": spec["status"],
        "supersedes": spec["supersedes"],
        "allowed_roles": spec["allowed_roles"],
        "topics": spec["topics"],
    }
    with open(sidecar_path, "w", encoding="utf-8") as f:
        json.dump(sidecar_data, f, indent=2, ensure_ascii=False)


def get_llm_prompt(spec: Dict[str, Any]) -> str:
    """Construct prompt for Groq client when generating rich prose."""
    source_type = spec["source_type"]
    prompt = (
        f"You are a technical pharmaceutical writer at Novacure Labs.\n"
        f"Draft realistic, highly detailed, GxP-compliant content for the following document:\n\n"
        f"Document ID: {spec['doc_id']}\n"
        f"Title: {spec['title']}\n"
        f"Document Type: {source_type}\n"
        f"Author: {spec['author']}\n"
        f"Date: {spec['date']}\n"
        f"Topics: {', '.join(spec['topics'])}\n\n"
        f"Content Brief:\n{spec['brief']}\n\n"
        f"MANDATORY REQUIREMENTS:\n"
    )
    if spec.get("fixed_sentences"):
        prompt += "1. The following sentences MUST appear VERBATIM in the text:\n"
        for s in spec["fixed_sentences"]:
            prompt += f'   - "{s}"\n'
    else:
        prompt += "1. Write authentic, professional GxP content.\n"

    prompt += (
        "2. The text MUST NOT contain the string 'HPLC' under any circumstances.\n"
        "3. Key people who must appear or be referenced: "
        + ", ".join(spec["people"]) + ".\n"
    )
    if spec["doc_id"] in ["REG-001", "EQP-001", "STB-001"]:
        prompt += "4. Do NOT state any temperature range other than '2 to 8 °C'.\n"

    if source_type == "sop":
        prompt += "5. Format as numbered sections: 1. Purpose, 2. Scope, 3. Responsibilities, 4. Procedure, etc.\n"
    elif source_type == "sharepoint":
        prompt += "5. Format with headings and paragraphs.\n"
    elif source_type == "email":
        prompt += "5. Format as an email with Subject line and body text.\n"
    elif source_type == "teams":
        prompt += "5. Format as 10 to 18 chat messages, one per line: [Author] (HH:MM): Message\n"

    return prompt


def generate_document(spec: Dict[str, Any], offline: bool = False, force: bool = False) -> bool:
    """Generate a single document and its sidecar."""
    doc_id = spec["doc_id"]
    folder = spec["folder"]
    filename = spec["filename"]
    target_dir = DATA_RAW_DIR / folder
    target_file = target_dir / filename
    sidecar_file = Path(f"{target_file}.meta.json")

    # Resumability check
    if not force and target_file.exists() and sidecar_file.exists():
        print(f"[SKIP] {doc_id} already exists at {target_file}")
        return True

    print(f"[GEN] Generating {doc_id} ({spec['title']})...")

    content = None
    if not offline:
        # Attempt LLM generation
        prompt = get_llm_prompt(spec)
        llm_response = llm_generate(prompt)
        time.sleep(2)  # Pause 2 seconds between LLM calls per spec
        if llm_response and len(llm_response.strip()) > 50:
            llm_clean = sanitize_text(llm_response)
            llm_with_fixed = ensure_fixed_sentences(llm_clean, spec.get("fixed_sentences", []))
            content = llm_with_fixed

    # Fallback to offline template if LLM returned empty string or offline mode
    if content is None:
        content = TEMPLATES.get(doc_id)
        if content is None:
            sys.stderr.write(f"No template found for {doc_id}\n")
            return False

    # Render according to source type
    st = spec["source_type"]
    if st in ["sop", "pdf"]:
        render_pdf(target_file, spec, content)
    elif st == "sharepoint":
        render_docx(target_file, spec, content)
    elif st == "email":
        render_eml(target_file, spec, content)
    elif st == "teams":
        render_teams(target_file, spec, content)
    else:
        sys.stderr.write(f"Unknown source type: {st}\n")
        return False

    # Write sidecar
    write_sidecar(target_file, spec)
    print(f"[OK] Written {target_file} and {sidecar_file}")
    return True


def main():
    parser = argparse.ArgumentParser(description="Novacure Labs Corpus Generator")
    parser.add_argument("--offline", action="store_true", help="Generate using built-in offline templates only")
    parser.add_argument("--force", action="store_true", help="Overwrite existing generated documents")
    parser.add_argument("--only", type=str, default=None, help="Generate only specified DOC_ID")
    args = parser.parse_args()

    # Ensure output directories exist
    for sub in ["sop", "pdf", "sharepoint", "email", "teams"]:
        (DATA_RAW_DIR / sub).mkdir(parents=True, exist_ok=True)

    target_specs = MANIFEST
    if args.only:
        target_specs = [s for s in MANIFEST if s["doc_id"] == args.only]
        if not target_specs:
            sys.stderr.write(f"Error: doc_id '{args.only}' not found in manifest.\n")
            sys.exit(1)

    print(f"Starting corpus generation (total specs: {len(target_specs)}, offline={args.offline}, force={args.force})...")
    success_count = 0
    for spec in target_specs:
        ok = generate_document(spec, offline=args.offline, force=args.force)
        if ok:
            success_count += 1
        else:
            sys.stderr.write(f"Failed to generate {spec['doc_id']}\n")

    print(f"Generation completed: {success_count}/{len(target_specs)} documents successfully created.")


if __name__ == "__main__":
    main()
